"""
Exporta el catálogo de resolución a un Excel navegable: data/catalogo/catalogo.xlsx.

Hojas:
  Leeme        qué contiene cada hoja y cómo se leen las dimensiones
  Resumen      cuadros y pares cuadro-serie por familia y vigencia (fórmulas sobre Cuadros/Series)
  Cuadros      una fila por cuadro: tupla del cuadro, eje de apertura y la cabeza que lo pide completo
  Series       una fila por par cuadro-serie con la tupla completa (14 dimensiones)
  Dimensiones  valores de cada dimensión con etiqueta, conteos (fórmulas) y si el modelo los ve

Uso:
    python util/export_catalog_xlsx.py [--out data/catalogo/catalogo.xlsx]
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from catalog import DIMS, VALORES_MARGINALES, Catalog  # noqa: E402

FONT = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1F3864")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT)
TITLE_FONT = Font(name=FONT, bold=True, size=14)
BOLD = Font(name=FONT, bold=True)

FAMILIAS = ["PIB regional por actividad económica", "Gasto del PIB", "PIB por actividad económica",
            "Imacec", "PIB por región", "PIB agregado"]
CABEZA = {"actividad": "activ_cls", "region": "reg_cls", "gasto": "gast_cls"}

# Etiquetas de las dimensiones de cuadro (las de categoría se toman del nombre de la serie en la BDE)
ETIQUETAS = {
    "indicador": {"pib": "Producto interno bruto", "imacec": "Indicador mensual de actividad económica"},
    "presentacion": {"nivel": "Nivel", "contribucion_aa": "Contribución respecto de igual período del año anterior",
                     "contribucion_pa": "Contribución respecto del período anterior",
                     "participacion": "Participación porcentual", "tasa": "Tasa (inversión / PIB)"},
    "valoracion": {"encadenado": "Volumen a precios del año anterior encadenado", "corrientes": "Precios corrientes",
                   "constantes": "Precios constantes de año base", "na": "No aplica (índice)"},
    "estacionalidad": {"original": "Serie original", "desestacionalizada": "Desestacionalizada",
                       "tendencia_ciclo": "Tendencia-ciclo"},
    "frecuencia": {"mensual": "Mensual", "trimestral": "Trimestral", "anual": "Anual"},
    "apertura": {"actividad": "Por actividad económica", "region": "Por región", "gasto": "Por componente del gasto",
                 "total": "Sin apertura"},
    "vigencia": {"vigente": "Publicación vigente", "historico": "Compilación o base superada"},
    "empalme": {"sí": "Series empalmadas", "no": "Sin empalme"},
    "info_historica": {"sí": "Información histórica", "no": "—"},
}
DESCRIPCION = {
    "indicador": "Indicador publicado.",
    "presentacion": "Forma de presentación del dato. Sin mención en la consulta: nivel (preferencia).",
    "valoracion": "Valoración del PIB. Sin mención: volumen encadenado (preferencia).",
    "referencia": "Compilación de referencia o base. Se resuelve fuera del modelo; sin mención: 2018.",
    "empalme": "Si las series están empalmadas. Preferencia de desempate: sí.",
    "info_historica": "Si el cuadro es información histórica. Constante: no.",
    "estacionalidad": "Tratamiento estacional. Sin mención: tabla condicionada (pendiente) o pregunta aclaratoria.",
    "frecuencia": "Frecuencia de la serie (un cuadro puede mezclar frecuencias).",
    "apertura": "Eje por el que se abre el cuadro; su cabeza pide el cuadro completo.",
    "region_fija": "Región de los cuadros región x actividad (una región por cuadro).",
    "vigencia": "Vigencia de la publicación. Constante: vigente.",
    "actividad": "Categoría de actividad económica (fila del cuadro).",
    "region": "Región (fila de los cuadros por región; en los cuadros región x actividad repite la región fija).",
    "gasto": "Componente del gasto (fila del cuadro).",
}
DIMS_CUADRO_XL = ["indicador", "presentacion", "valoracion", "referencia", "empalme", "info_historica",
                  "estacionalidad", "apertura", "region_fija", "vigencia"]


def familia(s: dict) -> str:
    if s["indicador"] == "imacec":
        return "Imacec"
    if s["region_fija"]:
        return "PIB regional por actividad económica"
    return {"gasto": "Gasto del PIB", "actividad": "PIB por actividad económica", "region": "PIB por región",
            "total": "PIB agregado"}[s["apertura"]]


def celda(v):
    if isinstance(v, bool):
        return "sí" if v else "no"
    return v


def header(ws, cols: list[str], widths: list[int]) -> None:
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "B2"


def body_font(ws) -> None:
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY_FONT


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--out", default="data/catalogo/catalogo.xlsx")
    a = p.parse_args()

    cat = Catalog.load()
    vocab = json.loads((Path(__file__).resolve().parents[1] / "data/catalogo/vocabularios.json").read_text())
    modelo = {k.lower(): set(v) for k, v in vocab["modelo"]["entidades"].items()}

    por_cuadro: dict[str, list[dict]] = defaultdict(list)
    for s in cat:
        por_cuadro[s["cuadro_id"]].append(s)

    wb = Workbook()
    leeme = wb.active
    leeme.title = "Leeme"
    resumen = wb.create_sheet("Resumen")
    wc = wb.create_sheet("Cuadros")
    ws = wb.create_sheet("Series")
    wd = wb.create_sheet("Dimensiones")

    # ---- Series ------------------------------------------------------------------------------
    cols_s = ["serie_id", "cuadro_id", "serie", "familia", "categoria"] + DIMS + ["unidad", "fecha_primer_valor"]
    header(ws, cols_s, [34, 9, 46, 30, 26] + [14] * len(DIMS) + [30, 14])
    for s in sorted(cat, key=lambda s: (s["cuadro_id"], s["serie_id"])):
        ws.append([s["serie_id"], s["cuadro_id"], s["serie"], familia(s), s["categoria"]] +
                  [celda(s[d]) for d in DIMS] + [s["unidad"], s["fecha_primer_valor"]])
    body_font(ws)
    ws.auto_filter.ref = ws.dimensions
    n_series = ws.max_row
    col_s = {c: get_column_letter(i) for i, c in enumerate(cols_s, 1)}

    # ---- Cuadros -----------------------------------------------------------------------------
    cols_c = ["cuadro_id", "cuadro", "familia"] + DIMS_CUADRO_XL + ["frecuencias", "cabeza_completa", "n_series",
                                                                    "categorias"]
    header(wc, cols_c, [9, 70, 30] + [14] * len(DIMS_CUADRO_XL) + [20, 16, 10, 60])
    i_n = cols_c.index("n_series") + 1
    for r, cid in enumerate(sorted(por_cuadro), start=2):
        rows = por_cuadro[cid]
        s = rows[0]
        frec = ", ".join(f for f in ("mensual", "trimestral", "anual") if any(x["frecuencia"] == f for x in rows))
        cabeza = ("activ_cls (con región fija)" if s["region_fija"] else CABEZA.get(s["apertura"], "—"))
        cats = sorted({str(x["categoria"]) for x in rows if x["categoria"]})
        wc.append([cid, s["cuadro"], familia(s)] + [celda(s[d]) for d in DIMS_CUADRO_XL] +
                  [frec, cabeza, None, ", ".join(cats)])
        wc.cell(row=r, column=i_n, value=f"=COUNTIF(Series!${col_s['cuadro_id']}$2:${col_s['cuadro_id']}${n_series},A{r})")
    body_font(wc)
    wc.auto_filter.ref = wc.dimensions
    n_cuadros = wc.max_row
    col_c = {c: get_column_letter(i) for i, c in enumerate(cols_c, 1)}

    # ---- Dimensiones -------------------------------------------------------------------------
    cols_d = ["dimension", "valor", "etiqueta", "en_vocabulario_del_modelo", "pares_total", "pares_vigentes",
              "descripcion_de_la_dimension"]
    header(wd, cols_d, [16, 32, 52, 16, 12, 14, 70])
    nombres = defaultdict(Counter)   # etiqueta de categoría: nombre de serie más frecuente
    for s in cat:
        for d in ("actividad", "region", "gasto"):
            if s[d] and s["categoria"] == s[d]:
                nombres[(d, s[d])][s["serie"]] += 1
    rs = f"$2:${n_series}"
    vig = f"Series!${col_s['vigencia']}${2}:${col_s['vigencia']}${n_series}"
    r = 2
    for d in DIMS:
        vals = Counter(celda(s[d]) for s in cat if s[d] is not None)
        for v in sorted(vals, key=lambda v: (-vals[v], str(v))):
            etiqueta = ETIQUETAS.get(d, {}).get(v)
            if etiqueta is None and (d, v) in nombres:
                etiqueta = nombres[(d, v)].most_common(1)[0][0]
            if d in modelo:
                en_modelo = "sí" if str(v) in modelo[d] else ("marginal" if v in VALORES_MARGINALES.get(d, ()) else "no")
            else:
                en_modelo = {"referencia": "regla fuera del modelo", "apertura": "derivada del catálogo",
                             "region_fija": "derivada del catálogo"}.get(d, "valor por defecto")
            rng = f"Series!${col_s[d]}${2}:${col_s[d]}${n_series}"
            wd.append([d, v, etiqueta, en_modelo, f"=COUNTIF({rng},B{r})",
                       f'=COUNTIFS({rng},B{r},{vig},"vigente")', DESCRIPCION.get(d, "")])
            r += 1
    body_font(wd)
    wd.auto_filter.ref = wd.dimensions

    # ---- Resumen -----------------------------------------------------------------------------
    resumen["A1"], resumen["A1"].font = "Universo del catálogo por familia y vigencia", TITLE_FONT
    cols_r = ["Familia", "Cuadros vigentes", "Cuadros históricos", "Cuadros total",
              "Pares vigentes", "Pares históricos", "Pares total"]
    resumen.append([])
    resumen.append(cols_r)
    for i, w in enumerate([40, 14, 14, 14, 14, 14, 14], 1):
        c = resumen.cell(row=3, column=i)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True)
        resumen.column_dimensions[get_column_letter(i)].width = w
    fc, vc = (f"Cuadros!${col_c[k]}$2:${col_c[k]}${n_cuadros}" for k in ("familia", "vigencia"))
    fs, vs = (f"Series!${col_s[k]}$2:${col_s[k]}${n_series}" for k in ("familia", "vigencia"))
    for i, fam in enumerate(FAMILIAS, start=4):
        resumen.append([fam, f'=COUNTIFS({fc},A{i},{vc},"vigente")', f'=COUNTIFS({fc},A{i},{vc},"historico")',
                        f"=B{i}+C{i}", f'=COUNTIFS({fs},A{i},{vs},"vigente")',
                        f'=COUNTIFS({fs},A{i},{vs},"historico")', f"=E{i}+F{i}"])
    last = 3 + len(FAMILIAS)
    resumen.append(["Total"] + [f"=SUM({get_column_letter(j)}4:{get_column_letter(j)}{last})" for j in range(2, 8)])
    for row in resumen.iter_rows(min_row=4, max_row=last + 1):
        for c in row:
            c.font = BOLD if c.row == last + 1 else BODY_FONT
            if c.column > 1:
                c.number_format = "#,##0"
    resumen.cell(row=last + 3, column=1,
                 value="Códigos de serie distintos (una serie puede figurar en más de un cuadro): "
                       f"{len({s['serie_id'] for s in cat}):,}".replace(",", ".")).font = BODY_FONT
    resumen.cell(row=last + 4, column=1, value="Fuente: exportación de la BDE de agosto de 2026, "
                 "data/catalogo/raw/catalogo_series_bde_2026-08-05.xlsx").font = BODY_FONT

    # ---- Leeme -------------------------------------------------------------------------------
    leeme.column_dimensions["A"].width = 18
    leeme.column_dimensions["B"].width = 110
    lineas = [
        ("Catálogo de resolución de PIBERT", None),
        (None, None),
        ("Hoja", "Contenido"),
        ("Resumen", "Cuadros y pares cuadro-serie por familia y vigencia (se recalcula desde Cuadros y Series)."),
        ("Cuadros", "Una fila por cuadro: su tupla, las frecuencias que contiene, la cabeza que lo pide completo "
                    "y sus categorías. Filtre por cualquier columna."),
        ("Series", "Una fila por par cuadro-serie con el código BDE y las 14 dimensiones de su tupla. "
                   "Una misma serie puede figurar en más de un cuadro (cortes alternativos del mismo dato)."),
        ("Dimensiones", "Valores de cada dimensión con su etiqueta, cuántos pares lo tienen y si el modelo lo ve "
                        "como entidad ('sí'), si es un valor marginal excluido, o si la dimensión se deriva del "
                        "catálogo o se resuelve fuera del modelo."),
        (None, None),
        ("Cómo se lee", None),
        ("Tupla", "Cada serie se describe por 14 dimensiones derivadas de su código oficial. Ocho son "
                  "enunciables en una consulta (indicador, frecuencia, estacionalidad, presentación, valoración, "
                  "actividad, región, gasto); las demás se fijan por defecto o se derivan del catálogo."),
        ("Cabezas", "activ_cls, reg_cls y gast_cls indican que la consulta pide la clasificación completa del eje. "
                    "Corresponden a la apertura del cuadro: actividad, región o gasto; los cuadros con región fija "
                    "son región x actividad."),
        ("Fila agregada", "categoria = total (o vacía) es la fila agregada del cuadro."),
        ("Vacío", "Una celda vacía significa que la dimensión no aplica a esa familia."),
        (None, None),
        ("Generado por", "util/export_catalog_xlsx.py a partir de data/catalogo/series.csv. No editar a mano: "
                         "regenerar con el script."),
    ]
    for i, (k, v) in enumerate(lineas, 1):
        leeme.cell(row=i, column=1, value=k).font = TITLE_FONT if i == 1 else BOLD
        c = leeme.cell(row=i, column=2, value=v)
        c.font = BOLD if i == 3 else BODY_FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
        leeme.cell(row=i, column=1).alignment = Alignment(vertical="top")

    out = Path(a.out)
    wb.save(out)
    print(f"{out}: {n_cuadros - 1} cuadros, {n_series - 1} pares cuadro-serie, {wd.max_row - 1} valores de dimensión")


if __name__ == "__main__":
    main()
